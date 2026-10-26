#!/usr/bin/env python3
"""CLASSFIND seed extractor.

Reads the real "Students List B.Tech I Year.xlsx" workbook and produces:
  - data/seed/students.generated.json  -> minimal student records
  - data/seed/import-report.txt        -> import statistics

Only the minimal fields required by the product are exported:
  name, mobile, program, section, rollNumber
"""
import json
import os
import re
import sys

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
SEED_DIR = os.path.join(ROOT, "data", "seed")

DEFAULT_SRC = os.path.join(
    os.path.expanduser("~"),
    "Downloads",
    "Telegram Desktop",
    "Students List B.Tech I Year.xlsx",
)

STUDENT_COLUMNS = {
    "Admission No": 1,
    "Program": 2,
    "Student Name": 5,
    "Mobile": 6,
    "Sec": 13,
    "Sec Roll No": 14,
}


def normalize_name(raw):
    if raw is None:
        return ""
    return re.sub(r"\s+", " ", str(raw)).strip()


def normalize_mobile(raw):
    """Normalize an Indian mobile number to a 10-digit string (or '')."""
    if raw is None:
        return ""
    digits = re.sub(r"\D", "", str(raw))
    if len(digits) == 12 and digits.startswith("91"):
        digits = digits[2:]
    elif len(digits) == 11 and digits.startswith("0"):
        digits = digits[1:]
    elif len(digits) > 10 and digits.startswith("91"):
        digits = digits[-10:]
    if len(digits) == 10:
        return digits
    return ""


def extract_sheet(ws):
    """Yield dicts for the student sheets with headers on row 3."""
    rows = ws.iter_rows(values_only=True)
    for _ in range(3):
        next(rows, None)  # skip title rows + header
    for row in rows:
        if row is None or len(row) < 15:
            continue
        name = normalize_name(row[5])
        mobile = normalize_mobile(row[6])
        sec = (str(row[13]).strip().upper() if row[13] else "")
        roll = (str(row[14]).strip() if row[14] else "")
        prog = (str(row[2]).strip().upper() if row[2] else "")
        if not name:
            continue
        yield {
            "name": name,
            "mobile": mobile,
            "program": prog,
            "section": sec,
            "rollNumber": roll,
        }


def main():
    src = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_SRC
    if not os.path.exists(src):
        print(f"Source not found: {src}", file=sys.stderr)
        sys.exit(1)

    wb = load_workbook(src, read_only=True, data_only=True)

    # raw rows extracted from the section sheets (may include bad mobiles)
    raw_rows = []
    for sheet in ("GROUP 1", "Group 2"):
        if sheet in wb.sheetnames:
            raw_rows.extend(extract_sheet(wb[sheet]))
    total = len(raw_rows)

    # -- validation -----------------------------------------------------
    seen = {}
    duplicate_pairs = 0
    invalid = 0
    no_section = 0
    valid = []
    for r in raw_rows:
        if not re.fullmatch(r"\d{10}", r["mobile"]):
            invalid += 1
            continue
        key = (r["mobile"], r["name"].lower())
        if key in seen:
            duplicate_pairs += 1
        else:
            seen[key] = True
        if not r["section"]:
            no_section += 1
        valid.append(r)

    os.makedirs(SEED_DIR, exist_ok=True)
    out_path = os.path.join(SEED_DIR, "students.generated.json")
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(valid, f, ensure_ascii=False, indent=1)

    sections = sorted({r["section"] for r in valid if r["section"]})
    report = (
        f"CLASSFIND IMPORT REPORT\n"
        f"======================\n"
        f"Source: {src}\n"
        f"Records read: {total}\n"
        f"Invalid mobile (excluded): {invalid}\n"
        f"Records exported: {len(valid)}\n"
        f"Duplicate pairs flagged: {duplicate_pairs}\n"
        f"Missing section: {no_section}\n"
        f"Sections detected: {len(sections)}\n"
        f"Sections: {', '.join(sections)}\n"
        f"Output: {out_path}\n"
    )
    report_path = os.path.join(SEED_DIR, "import-report.txt")
    with open(report_path, "w", encoding="utf-8") as f:
        f.write(report)

    print(report)
    print("OK — seed written.")


if __name__ == "__main__":
    main()